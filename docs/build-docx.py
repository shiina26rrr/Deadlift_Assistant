from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
SRC = ROOT / "work-report.md"
OUT = ROOT / "作品说明文档-v1.1.1.docx"

TITLE = "硬拉辅助检测系统"
SUBTITLE = "作品说明文档"
META = [
    ("文档版本", "1.0"),
    ("应用版本", "1.1.1"),
    ("日期", "2026-10-26"),
    ("编写方", "Deadlift Assistant Team"),
    ("适用范围", "Windows 10/11 x64、传统杠铃硬拉、初学者练习辅助"),
]

BODY_FONT = "宋体"
HEAD_FONT = "黑体"
CODE_FONT = "Consolas"


def set_run_font(run, name: str, size: float | None = None, bold: bool | None = None,
                 color: RGBColor | None = None, italic: bool | None = None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def shade(element, fill: str):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    element.append(shd)


def add_field(paragraph, instr: str, placeholder: str, size: float = 9):
    run = paragraph.add_run()
    set_run_font(run, BODY_FONT, size)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr_el = OxmlElement("w:instrText")
    instr_el.set(qn("xml:space"), "preserve")
    instr_el.text = instr
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for el in (begin, instr_el, sep, text, end):
        run._r.append(el)


INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")


def add_inline(paragraph, text: str, size: float = 10.5, base_bold: bool = False):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, BODY_FONT, size, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, CODE_FONT, size - 0.5)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, BODY_FONT, size, bold=base_bold)


def configure_styles(doc: Document):
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(4)

    for name, size, before, after in (
        ("Heading 1", 18, 18, 12),
        ("Heading 2", 14, 14, 8),
        ("Heading 3", 11.5, 10, 6),
    ):
        style = doc.styles[name]
        style.font.name = HEAD_FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), HEAD_FONT)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def page_setup(section):
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.4)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.2)
    section.header_distance = Cm(1.2)
    section.footer_distance = Cm(1.2)


def build_cover(doc: Document):
    for _ in range(5):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(TITLE)
    set_run_font(run, HEAD_FONT, 30, bold=True, color=RGBColor(0x1F, 0x3A, 0x5F))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(24)
    run = p.add_run(SUBTITLE)
    set_run_font(run, HEAD_FONT, 22, bold=True, color=RGBColor(0x33, 0x33, 0x33))

    line = doc.add_paragraph()
    line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pPr = line._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1F3A5F")
    borders.append(bottom)
    pPr.append(borders)

    doc.add_paragraph()
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for key, value in META:
        cells = table.add_row().cells
        cells[0].width = Cm(3.6)
        cells[1].width = Cm(9.6)
        p0 = cells[0].paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r0 = p0.add_run(f"{key}：")
        set_run_font(r0, HEAD_FONT, 11, bold=True)
        p1 = cells[1].paragraphs[0]
        r1 = p1.add_run(value)
        set_run_font(r1, BODY_FONT, 11)

    doc.add_paragraph()
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("本文件为作品说明与验证状态记录，不构成训练处方、医学建议或安全结论。")
    set_run_font(run, BODY_FONT, 10, color=RGBColor(0x66, 0x66, 0x66), italic=True)

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def build_toc(doc: Document):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(16)
    run = p.add_run("目  录")
    set_run_font(run, HEAD_FONT, 20, bold=True, color=RGBColor(0x1F, 0x3A, 0x5F))

    toc = doc.add_paragraph()
    add_field(toc, 'TOC \\o "1-2" \\h \\z \\u', "打开 Word 后按 Ctrl+A 再按 F9，或右键“更新域”生成目录。", 10.5)

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(12)
    add_inline(note, "提示：目录为 Word 域，页码随排版自动更新；若以 PDF 形式分发，导出前已在 Word 中更新域。", 9.5)

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_table(doc: Document, rows: list[list[str]]):
    header = rows[0]
    body = rows[1:]
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    hdr = table.rows[0]
    trPr = hdr._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)

    for i, text in enumerate(header):
        cell = hdr.cells[i]
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(0)
        add_inline(p, text, 9, base_bold=True)
        shade(cell._tc.get_or_add_tcPr(), "E8EEF5")

    for row in body:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            p = cells[i].paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(0)
            add_inline(p, text, 9)

    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_code_block(doc: Document, lines: list[str]):
    for line in lines:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.5)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(0)
        pPr = p._p.get_or_add_pPr()
        shade(pPr, "F4F4F4")
        run = p.add_run(line if line else " ")
        set_run_font(run, CODE_FONT, 9)


def add_quote(doc: Document, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.6)
    p.paragraph_format.line_spacing = 1.3
    add_inline(p, text, 9.5)
    for run in p.runs:
        run.italic = True
        run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    pPr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "12")
    left.set(qn("w:space"), "6")
    left.set(qn("w:color"), "1F3A5F")
    borders.append(left)
    pPr.append(borders)


def parse_table(lines: list[str]) -> list[list[str]]:
    rows = []
    for line in lines:
        stripped = line.strip().strip("|")
        cells = [c.strip() for c in stripped.split("|")]
        rows.append(cells)
    return [rows[0]] + rows[2:]


def render_body(doc: Document, lines: list[str]):
    i = 0
    total = len(lines)
    while i < total:
        raw = lines[i]
        line = raw.rstrip()
        if not line.strip():
            i += 1
            continue
        if line.startswith("```"):
            block = []
            i += 1
            while i < total and not lines[i].startswith("```"):
                block.append(lines[i].rstrip())
                i += 1
            i += 1
            add_code_block(doc, block)
            continue
        if line.startswith("|"):
            block = []
            while i < total and lines[i].lstrip().startswith("|"):
                block.append(lines[i])
                i += 1
            add_table(doc, parse_table(block))
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=3)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=1)
        elif line.startswith("> "):
            add_quote(doc, line[2:].strip())
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, line[2:].strip())
        elif line.strip() == "---":
            pass
        else:
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Cm(0.74)
            add_inline(p, line.strip())
        i += 1


def add_header_footer(section):
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("硬拉辅助检测系统 作品说明文档 · v1.0")
    set_run_font(run, BODY_FONT, 9, color=RGBColor(0x66, 0x66, 0x66))

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("第 ")
    set_run_font(r, BODY_FONT, 9)
    add_field(p, "PAGE", "1", 9)
    r = p.add_run(" 页 / 共 ")
    set_run_font(r, BODY_FONT, 9)
    add_field(p, "NUMPAGES", "1", 9)
    r = p.add_run(" 页")
    set_run_font(r, BODY_FONT, 9)


def main():
    text = SRC.read_text(encoding="utf-8")
    lines = text.splitlines()

    start = next(i for i, l in enumerate(lines) if l.strip() == "## 免责声明")
    toc_start = next(i for i, l in enumerate(lines) if l.strip() == "## 目录")
    ch1 = next(i for i, l in enumerate(lines) if l.startswith("# 第一章"))

    doc = Document()
    configure_styles(doc)
    page_setup(doc.sections[0])
    build_cover(doc)
    build_toc(doc)

    body = lines[start:toc_start] + lines[ch1:]
    render_body(doc, body)

    add_header_footer(doc.sections[0])
    doc.sections[0].different_first_page_header_footer = True

    settings = doc.settings.element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)

    doc.core_properties.title = f"{TITLE} {SUBTITLE}"
    doc.core_properties.author = "Deadlift Assistant Team"
    doc.core_properties.subject = "应用版本 1.1.1 / 文档版本 1.0"
    doc.core_properties.comments = "作品说明文档；未实测项见附录 B。"

    doc.save(str(OUT))
    print(f"Saved: {OUT} ({OUT.stat().st_size} bytes)")


if __name__ == "__main__":
    sys.exit(main())
